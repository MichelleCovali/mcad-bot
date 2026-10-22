"""Lossless, repeatable import of the private MCAD source workbook.

No workbook formula is evaluated and no geocoding or network request is made.
"""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime, time
from hashlib import sha256
import json
import math
from pathlib import Path
import posixpath
import re
import unicodedata
import warnings
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries


SOURCE_FIELDS = {
    "QA": "qa",
    "Reference number": "reference_number",
    "Month": "month",
    "Year": "year",
    "Year End (blank if same year)": "year_end",
    "Title": "title",
    "Impact area": "impact_area",
    "Incident Location": "location_name",
    "Incident Country": "incident_country",
    "Latitude Longitude": "coordinates_raw",
    "Google Map Location": "google_map_location",
    "Victim Country": "victim_country",
    "Victim identity": "victim_identity",
    "Victim type": "victim_type",
    "Victim type short": "victim_type_short",
    "Affected system": "affected_system",
    "Method": "method",
    "Social engineering technique": "social_engineering_technique",
    "Malware type": "malware_type",
    "Ransomware type": "ransomware_type",
    "Attack Pattern": "attack_pattern",
    "Motive": "motive",
    "Motive short": "motive_short",
    "Campaign": "campaign",
    "Impact": "impact",
    "Impact short": "impact_short",
    "Threat Actor country": "attacker_country",
    "Countermeasures": "countermeasures",
    "Threat Actor": "threat_actor",
    "Additional info": "additional_info",
    "Summary": "summary",
    "Reference short": "reference_short",
    "Reference URL": "reference_urls_raw",
}

COUNTRY_ALIASES = {
    "usa": "US", "us": "US", "u.s.": "US", "the us": "US",
    "uk": "GB", "u.k.": "GB", "uae": "AE",
    "the netherlands": "NL", "hong kong": "HK", "czech republic": "CZ",
    "macedonia": "MK", "libanon": "LB", "ukarine": "UA", "vietman": "VN",
    "myanmar": "MM",
}
UNKNOWN_COUNTRIES = {"unknown", "uknown", "n/a", "n.a.", "na", "other", "?", "-"}
MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NUMBER = r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)"


def _text(value):
    return "" if value is None else str(value).strip()


def _json_value(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    return value


def _identity_text(value):
    return " ".join(unicodedata.normalize("NFKC", _text(value)).casefold().split())


def external_identity(reference, title):
    """Titles disambiguate the source's repeated reference numbers, without row order."""
    digest = sha256(_identity_text(title).encode()).hexdigest()[:16]
    return f"mcad-excel:{_text(reference) or 'unreferenced'}:{digest}"


def parse_coordinates(value):
    """Accept only an explicit source pair; never invent or swap coordinates."""
    text = _text(value)
    if not text:
        return {"status": "missing", "latitude": None, "longitude": None, "separator": None}
    match = re.fullmatch(rf"\s*({NUMBER})\s*,\s*({NUMBER})\s*", text)
    separator = "comma"
    if not match:
        match = re.fullmatch(rf"\s*({NUMBER})\s+({NUMBER})\s*", text)
        separator = "whitespace"
    if match:
        latitude, longitude = float(match[1]), float(match[2])
        if math.isfinite(latitude) and math.isfinite(longitude) and -90 <= latitude <= 90 and -180 <= longitude <= 180:
            return {"status": "valid", "latitude": latitude, "longitude": longitude, "separator": separator}
    return {"status": "invalid", "latitude": None, "longitude": None, "separator": None}


def _integer(value):
    if isinstance(value, bool):
        return None
    if isinstance(value, int) or isinstance(value, float) and math.isfinite(value) and value.is_integer():
        return int(value)
    return int(_text(value)) if re.fullmatch(r"\d+", _text(value)) else None


def parse_date_fields(year, month=None, year_end=None):
    """Keep month/year precision; reference-number suffixes are never treated as days."""
    issues = []
    parsed_year, parsed_month, parsed_end = _integer(year), _integer(month), _integer(year_end)
    if parsed_year is None or not 1900 <= parsed_year <= datetime.now().year:
        if _text(year):
            issues.append("invalid_year")
        parsed_year = None
    if parsed_month is not None and not 1 <= parsed_month <= 12 or _text(month) and parsed_month is None:
        issues.append("invalid_month")
        parsed_month = None
    if parsed_end is not None and (parsed_year is None or parsed_end < parsed_year or parsed_end > datetime.now().year) or _text(year_end) and parsed_end is None:
        issues.append("invalid_year_end")
        parsed_end = None
    label = f"{parsed_year:04d}" if parsed_year is not None else ""
    if label and parsed_month is not None:
        label += f"-{parsed_month:02d}"
    precision = "month" if parsed_year is not None and parsed_month is not None else "year" if parsed_year is not None else "unknown"
    if parsed_end is not None and parsed_end != parsed_year:
        label += f"–{parsed_end:04d}"
        precision = "range"
    return {
        "year": parsed_year, "month": str(parsed_month) if parsed_month is not None else "",
        "year_end": parsed_end, "date_text": label or None,
        "date_start": None, "date_precision": precision, "issues": issues,
    }


def resolve_countries(value, country_catalog):
    """Resolve explicit names/known aliases against the workbook's own code catalog."""
    by_code = {item["id"].casefold(): item for item in country_catalog if item.get("id")}
    by_name = {item["value"].casefold(): item for item in country_catalog if item.get("id")}
    relations = []
    for part in _text(value).split(","):
        token = part.strip()
        if not token:
            continue
        lookup = re.sub(r"^and\s+", "", token, flags=re.I).casefold()
        item = by_name.get(lookup) or by_code.get(lookup) or by_code.get(COUNTRY_ALIASES.get(lookup, "").casefold())
        if item:
            relations.append({"id": item["id"], "name": item["value"], "source": token, "status": "resolved"})
        else:
            relations.append({"id": None, "name": token, "source": token, "status": "unknown" if lookup in UNKNOWN_COUNTRIES else "unresolved"})
    return relations


def reference_urls(value, hyperlinks=()):
    candidates = re.findall(r"https?://[^\s<>\"\u201c\u201d]+", _text(value))
    candidates.extend(link["target"] for link in hyperlinks if link.get("target", "").startswith(("https://", "http://")))
    result = []
    for candidate in candidates:
        candidate = candidate.rstrip(".,;")
        if candidate not in result:
            result.append(candidate)
    return result


def _relationships(archive, part):
    rel_part = posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")
    if rel_part not in archive.namelist():
        return {}
    return {item.attrib["Id"]: dict(item.attrib) for item in ET.fromstring(archive.read(rel_part))}


def _resolve_part(part, target):
    return target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(posixpath.dirname(part), target))


def _native_metadata(path):
    result = {}
    with zipfile.ZipFile(path) as archive:
        workbook_part = "xl/workbook.xml"
        relations = _relationships(archive, workbook_part)
        workbook = ET.fromstring(archive.read(workbook_part))
        for sheet in workbook.findall(f"{{{MAIN_NS}}}sheets/{{{MAIN_NS}}}sheet"):
            target = relations[sheet.attrib[f"{{{REL_NS}}}id"]]["Target"]
            part = _resolve_part(workbook_part, target)
            xml = ET.fromstring(archive.read(part))
            sheet_relations = _relationships(archive, part)
            links = []
            for element in xml.findall(f"{{{MAIN_NS}}}hyperlinks/{{{MAIN_NS}}}hyperlink"):
                link = dict(element.attrib)
                rel_id = link.pop(f"{{{REL_NS}}}id", None)
                link["target"] = sheet_relations.get(rel_id, {}).get("Target", "")
                link["target_mode"] = sheet_relations.get(rel_id, {}).get("TargetMode")
                links.append(link)
            tables = []
            for element in xml.findall(f"{{{MAIN_NS}}}tableParts/{{{MAIN_NS}}}tablePart"):
                relation = sheet_relations[element.attrib[f"{{{REL_NS}}}id"]]
                table_xml = ET.fromstring(archive.read(_resolve_part(part, relation["Target"])))
                tables.append({"attributes": dict(table_xml.attrib), "columns": [dict(column.attrib) for column in table_xml.findall(f"{{{MAIN_NS}}}tableColumns/{{{MAIN_NS}}}tableColumn")]})
            result[sheet.attrib["name"]] = {
                "hyperlinks": links, "tables": tables,
                "auto_filters": [dict(element.attrib) for element in xml.iter() if element.tag.rsplit("}", 1)[-1] == "autoFilter"],
                "data_validation_xml": [ET.tostring(element, encoding="unicode") for element in xml.iter() if element.tag.rsplit("}", 1)[-1] == "dataValidation"],
                "merged_cells": [element.attrib["ref"] for element in xml.findall(f"{{{MAIN_NS}}}mergeCells/{{{MAIN_NS}}}mergeCell")],
            }
    return result


def _cell_links(address, native):
    links = []
    for link in native.get("hyperlinks", []):
        reference = link.get("ref", "")
        if reference == address:
            links.append(link)
        elif ":" in reference:
            low_col, low_row, high_col, high_row = range_boundaries(reference)
            col, row, _, _ = range_boundaries(address)
            if low_col <= col <= high_col and low_row <= row <= high_row:
                links.append(link)
    return links


def _raw_row(cells, headers, cached_values, native, row_number):
    raw, source_cells = {}, {}
    for index, cell in enumerate(cells):
        address = f"{get_column_letter(index + 1)}{row_number}"
        header = headers[index] if index < len(headers) else None
        key = str(header) if header is not None else f"column_{get_column_letter(index + 1)}"
        if key in raw:
            key += f" [column {get_column_letter(index + 1)}]"
        value = cell.value
        if header is not None or value is not None:
            raw[key] = _json_value(value)
        source_cell = {"header": _json_value(header), "value": _json_value(value), "excel_type": getattr(cell, "data_type", "n"), "number_format": getattr(cell, "number_format", "General")}
        if getattr(cell, "data_type", None) == "f":
            source_cell.update(formula=value, cached_value=_json_value(cached_values[index]))
        links = _cell_links(address, native)
        if links:
            source_cell["hyperlinks"] = links
        source_cells[address] = source_cell
    raw["_cells"] = source_cells
    return raw


@dataclass
class ParsedRecord:
    source_row: int
    external_id: str
    payload: dict
    issues: list[str]


@dataclass
class ParsedWorkbook:
    source_name: str
    source_sha256: str
    row_count: int
    records: list[ParsedRecord]
    quarantined: list[dict]
    metadata_json: dict

    def report(self):
        mapped = sum(record.payload["latitude"] is not None for record in self.records)
        return {
            "source_name": self.source_name, "source_sha256": self.source_sha256,
            "source_rows": self.row_count, "accepted_rows": len(self.records),
            "quarantined_rows": len(self.quarantined), "mapped_rows": mapped,
            "unmapped_rows": len(self.records) - mapped,
            "accounted_rows": len(self.records) + len(self.quarantined),
            "warning_counts": dict(Counter(issue for record in self.records for issue in record.issues)),
            "row_warnings": [{"row": record.source_row, "reference": record.payload["reference_number"], "issues": record.issues} for record in self.records if record.issues],
            "quarantine": [{"source_row": item["source_row"], "reason": item["reason"]} for item in self.quarantined],
            "historical_notifications": False,
        }


def _parse_workbook(path):
    path = Path(path)
    source_sha = sha256(path.read_bytes()).hexdigest()
    native = _native_metadata(path)
    # The reader warns about unsupported editing features. We never save the workbook;
    # native validation/table/link definitions are preserved separately from the reader.
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Data Validation extension is not supported and will be removed")
        workbook = load_workbook(path, read_only=True, data_only=False)
        cached = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Database" not in workbook.sheetnames or "lists" not in workbook.sheetnames:
            raise ValueError("The workbook must contain Database and lists worksheets.")
        list_rows = list(workbook["lists"].iter_rows())
        cached_lists = list(cached["lists"].iter_rows(values_only=True))
        list_headers = [cell.value for cell in list_rows[0]]
        catalog_rows = []
        catalogs = {"countries": [], "impact_areas": [], "methods": [], "victim_types": []}
        for index, cells in enumerate(list_rows[1:], 1):
            if not any(_text(cell.value) for cell in cells):
                continue
            values = [cell.value for cell in cells]
            row = index + 1
            catalog_rows.append({"source_row": row, "raw_fields": _raw_row(cells, list_headers, cached_lists[index], native["lists"], row)})
            if _text(values[0]) and _text(values[1]):
                catalogs["countries"].append({"id": _text(values[0]), "value": _text(values[1]), "source_row": row})
            for column, name in [(4, "impact_areas"), (8, "methods"), (11, "victim_types")]:
                if _text(values[column]):
                    catalogs[name].append({"value": _text(values[column]), "source_row": row})

        sheet = workbook["Database"]
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        normalized_headers = [_text(header) for header in headers]
        missing_headers = set(SOURCE_FIELDS) - set(normalized_headers)
        if missing_headers:
            raise ValueError("The source workbook is missing columns: " + ", ".join(sorted(missing_headers)))
        cached_rows = list(cached["Database"].iter_rows(values_only=True))
        records, quarantined, row_count = [], [], 0
        reference_rows = defaultdict(list)
        for row_number, cells in enumerate(sheet.iter_rows(min_row=2), 2):
            if not any(_text(cell.value) for cell in cells):
                continue
            row_count += 1
            cached_values = cached_rows[row_number - 1]
            raw = _raw_row(cells, headers, cached_values, native["Database"], row_number)
            canonical = {SOURCE_FIELDS[header]: cells[index].value for index, header in enumerate(normalized_headers) if header in SOURCE_FIELDS}
            issues = []
            dates = parse_date_fields(canonical.get("year"), canonical.get("month"), canonical.get("year_end"))
            issues.extend(dates.pop("issues"))
            coordinate = parse_coordinates(canonical.get("coordinates_raw"))
            if coordinate["status"] != "valid":
                issues.append("coordinates_" + coordinate["status"])
            elif coordinate["separator"] == "whitespace":
                issues.append("coordinate_whitespace_separator")

            countries = {}
            for field in ["incident_country", "victim_country", "attacker_country"]:
                countries[field] = resolve_countries(canonical.get(field), catalogs["countries"])
                if any(item["status"] == "unresolved" for item in countries[field]):
                    issues.append(field + "_unresolved")
            raw["_country_relations"] = countries
            raw["_source"] = {"workbook": path.name, "sha256": source_sha, "sheet": "Database", "row": row_number, "coordinate_parse": coordinate}

            title = _text(canonical.get("title"))
            if not title:
                quarantined.append({"source_row": row_number, "reason": "Missing Title; source row retained for review.", "raw_fields": raw})
                continue
            payload = {key: _text(value) for key, value in canonical.items() if key not in {"coordinates_raw", "reference_urls_raw", "year", "month", "year_end"}}
            payload.update(dates)
            payload.update(
                title=title, description=_text(canonical.get("summary")), status="published",
                latitude=coordinate["latitude"], longitude=coordinate["longitude"],
                location_precision="source_provided" if coordinate["status"] == "valid" else "unknown",
                location_source=f"{path.name}:Database!J{row_number}",
                raw_fields=raw,
            )
            for field, relations in countries.items():
                payload[field] = ", ".join(dict.fromkeys(item["name"] for item in relations))
            map_column = normalized_headers.index("Google Map Location")
            if getattr(cells[map_column], "data_type", None) == "f":
                payload["google_map_location"] = _text(cached_values[map_column])
            url_column = normalized_headers.index("Reference URL")
            urls = reference_urls(canonical.get("reference_urls_raw"), _cell_links(f"{get_column_letter(url_column + 1)}{row_number}", native["Database"]))
            payload.update(
                reference_urls=urls, source_url=urls[0] if urls else None,
                sources=[{"url": url, "short_name": _text(canonical.get("reference_short")), "raw_text": _text(canonical.get("reference_urls_raw"))} for url in urls],
                organisations=[{"name": payload["victim_identity"], "country": payload["victim_country"], "role": "victim"}] if payload["victim_identity"] else [],
                actors=[{"name": payload["threat_actor"], "country": payload["attacker_country"]}] if payload["threat_actor"] else [],
                locations=[{
                    "name": payload["location_name"], "country": payload["incident_country"],
                    "latitude": payload["latitude"], "longitude": payload["longitude"],
                    "precision": payload["location_precision"], "source": payload["location_source"],
                }],
            )
            reference = payload["reference_number"]
            reference_rows[reference].append(row_number)
            if re.fullmatch(r"\d{8}", reference):
                if dates["year"] is not None and int(reference[:4]) != dates["year"]:
                    issues.append("reference_year_mismatch")
                if int(reference[4:6]) != (_integer(dates["month"]) or 0):
                    issues.append("reference_month_mismatch")
            records.append(ParsedRecord(row_number, external_identity(reference, title), payload, issues))

        by_identity = defaultdict(list)
        for record in records:
            by_identity[record.external_id].append(record)
            if len(reference_rows[record.payload["reference_number"]]) > 1:
                record.issues.append("duplicate_reference_preserved")
        for group in by_identity.values():
            if len(group) < 2:
                continue
            repeated = Counter()
            for record in group:
                source_values = {key: value for key, value in record.payload["raw_fields"].items() if not key.startswith("_")}
                digest = sha256(json.dumps(source_values, sort_keys=True, ensure_ascii=False, default=_json_value).encode()).hexdigest()[:16]
                repeated[digest] += 1
                record.external_id += f":{digest}:occurrence:{repeated[digest]}"
                record.issues.append("duplicate_reference_title_preserved")

        metadata = {
            "format_version": 1, "sheet_names": workbook.sheetnames,
            "source_headers": headers, "source_field_mapping": SOURCE_FIELDS,
            "source_sheet_dimensions": {worksheet.title: worksheet.calculate_dimension() for worksheet in workbook.worksheets},
            "catalogs": catalogs, "lists_rows": catalog_rows,
            "defined_names": [{"name": name, "formula": value.attr_text} for name, value in workbook.defined_names.items()],
            "native_metadata": native,
            "duplicate_references": {key: rows for key, rows in reference_rows.items() if len(rows) > 1},
            "quarantined_rows": quarantined,
            "identity_policy": "reference + normalized-title SHA256; edited titles require manual identity reconciliation",
            "location_policy": "Source coordinates only, with unspecified source accuracy; no network geocoding or invented country centroids.",
            "notification_policy": "Historical import never emits notification events.",
        }
        parsed = ParsedWorkbook(path.name, source_sha, row_count, records, quarantined, metadata)
        if len(records) + len(quarantined) != row_count:
            raise RuntimeError("Import accounting failed: not every nonblank source row is represented.")
        return parsed
    finally:
        workbook.close()
        cached.close()


def parse_workbook(path):
    """Read every nonblank source row, retaining native metadata without saving XLSX."""
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Data Validation extension is not supported and will be removed")
        return _parse_workbook(path)


def _save_report(report, report_path):
    if report_path is not None:
        path = Path(report_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")


def import_workbook(path, session_factory=None, report_path=None):
    """Import atomically, suppress historical alerts, and skip a previously imported SHA."""
    from sqlalchemy import select
    from .db import get_session_factory
    from .models import ImportBatch, ImportRow
    from .service import upsert_incident

    parsed = parse_workbook(path)
    factory = session_factory or get_session_factory()
    report = parsed.report()
    with factory() as session:
        with session.begin():
            existing = session.scalar(select(ImportBatch).where(ImportBatch.source_sha256 == parsed.source_sha256).order_by(ImportBatch.created_at).limit(1))
            if existing:
                report.update(import_batch_id=existing.id, reused_batch=True, created_records=0, existing_records=len(parsed.records), notification_events_created=0)
            else:
                batch = ImportBatch(source_name=parsed.source_name, source_sha256=parsed.source_sha256, row_count=parsed.row_count, metadata_json=parsed.metadata_json)
                session.add(batch)
                session.flush()
                created_count = 0
                for record in parsed.records:
                    _, created = upsert_incident(session, record.payload, external_id=record.external_id, notify=False, import_batch_id=batch.id, source_row=record.source_row)
                    created_count += int(created)
                for row in parsed.quarantined:
                    session.add(ImportRow(import_batch_id=batch.id, source_row=row["source_row"], incident_id=None, raw_fields=row["raw_fields"]))
                session.flush()
                report.update(import_batch_id=batch.id, reused_batch=False, created_records=created_count, existing_records=len(parsed.records) - created_count, notification_events_created=0)
    _save_report(report, report_path)
    return report


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("path", type=Path, help="Private MCAD.xlsx source; never modified")
    parser.add_argument("--report", type=Path, default=Path("data/import-report.json"))
    parser.add_argument("--dry-run", action="store_true", help="Validate and report without opening a database")
    args = parser.parse_args(argv)
    if args.dry_run:
        report = parse_workbook(args.path).report()
        report.update(dry_run=True, created_records=0, notification_events_created=0)
        _save_report(report, args.report)
    else:
        report = import_workbook(args.path, report_path=args.report)
    summary_keys = ["source_name", "source_rows", "accepted_rows", "quarantined_rows", "mapped_rows", "unmapped_rows", "created_records", "existing_records", "reused_batch", "notification_events_created"]
    print(json.dumps({key: report[key] for key in summary_keys if key in report}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
